"""文本抽取（T-MAT · md/txt/pdf/docx/img/link）。

本地适配：
- 图片 OCR 若 PaddleOCR 不可用，优雅降级为占位文本（不阻断管线）。
- pdf 无文本层时尝试 OCR，同样可降级。
- ``raw_text:`` 前缀表示调用方直接传入文本（前端粘贴解析场景）。
"""
import asyncio
from pathlib import Path

RAW_PREFIX = "raw_text:"


async def extract_text(source: str, kind: str) -> str:
    if source.startswith(RAW_PREFIX):
        return source[len(RAW_PREFIX):]
    if kind in ("md", "txt"):
        return await asyncio.to_thread(_read_text, source)
    if kind == "pdf":
        return await asyncio.to_thread(_pdf, source)
    if kind == "docx":
        return await asyncio.to_thread(_docx, source)
    if kind in ("png", "jpg", "jpeg"):
        return await asyncio.to_thread(_ocr, source)
    if kind == "link":
        return await _link(source)
    # 兜底：当作纯文本路径
    return await asyncio.to_thread(_read_text, source)


def _read_text(path: str) -> str:
    return Path(path).read_text(encoding="utf-8", errors="ignore")


def _pdf(path: str) -> str:
    try:
        from pdfminer.high_level import extract_text as x

        txt = x(path)
        if txt and txt.strip():
            return txt
    except Exception:
        pass
    return _ocr(path)


def _docx(path: str) -> str:
    try:
        import docx

        return "\n".join(p.text for p in docx.Document(path).paragraphs)
    except Exception as e:  # pragma: no cover
        return f"[docx 解析失败: {e}]"


_engine = None
_engine_failed = False


def _ocr(path: str) -> str:
    """图片 OCR；PaddleOCR 不可用时降级占位（保证管线不崩）。"""
    global _engine, _engine_failed
    if _engine_failed:
        return f"[图片资料 OCR 不可用，已降级占位] {Path(path).name}"
    try:
        if _engine is None:
            from paddleocr import PaddleOCR

            _engine = PaddleOCR(use_angle_cls=True, lang="ch")
        res = _engine.ocr(path, cls=True)
        return "\n".join(line[1][0] for pg in res for line in pg)
    except Exception:
        _engine_failed = True
        return f"[图片资料 OCR 不可用，已降级占位] {Path(path).name}"


async def _link(url: str) -> str:
    import httpx

    async with httpx.AsyncClient(timeout=30, follow_redirects=True) as c:
        html = (await c.get(url)).text
    try:
        from readability import Document

        return Document(html).summary()
    except Exception:
        # 退化：去标签
        import re

        return re.sub(r"<[^>]+>", " ", html)
